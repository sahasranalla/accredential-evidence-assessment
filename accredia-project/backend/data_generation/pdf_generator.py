import os
import sys
import json
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas
import docx

INSTITUTION_NAME = "Vidyapath Institute of Technology"
INSTITUTION_SHORT = "VITEC"
LOCATION = "Hyderabad, Telangana, India"

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ACCREDITATION_DIR = os.path.join(BASE_DIR, "data", "accreditation")
DOCS_DIR = os.path.join(ACCREDITATION_DIR, "documents")
METADATA_DIR = os.path.join(ACCREDITATION_DIR, "metadata")
TEST_QUERIES_DIR = os.path.join(ACCREDITATION_DIR, "test_queries")

CRITERIA = [
    {
        "id": "Criterion 1",
        "name": "Curricular Aspects",
        "weightage": 100,
        "description": "Curriculum design, academic flexibility, value-added courses, and structured stakeholder feedback mechanisms.",
        "subcriteria": [
            {"id": "1.1", "name": "Curricular Planning and Implementation"},
            {"id": "1.2", "name": "Academic Flexibility and CBCS Electives"},
            {"id": "1.3", "name": "Curriculum Enrichment & Value-Added Courses"},
            {"id": "1.4", "name": "Feedback System on Institutional Curriculum"}
        ]
    },
    {
        "id": "Criterion 2",
        "name": "Teaching-Learning and Evaluation",
        "weightage": 350,
        "description": "Student diversity catering, experiential teaching methods, qualified faculty cadre, continuous internal evaluation, and CO-PO attainment.",
        "subcriteria": [
            {"id": "2.1", "name": "Student Enrollment & Reservation Profile"},
            {"id": "2.2", "name": "Catering to Diverse Learning Needs"},
            {"id": "2.3", "name": "Teaching-Learning Process & Experiential Pedagogy"},
            {"id": "2.4", "name": "Teacher Profile & Qualification Quality"},
            {"id": "2.5", "name": "Evaluation Process and Examination Reforms"},
            {"id": "2.6", "name": "Student Performance and Learning Outcomes (CO-PO)"},
            {"id": "2.7", "name": "Student Satisfaction Survey (SSS)"}
        ]
    },
    {
        "id": "Criterion 3",
        "name": "Research, Innovations and Extension",
        "weightage": 110,
        "description": "Resource mobilization for research, incubation ecosystem, Scopus publications, patents, community extension, and corporate collaborations.",
        "subcriteria": [
            {"id": "3.1", "name": "Resource Mobilization for Research & Seed Grants"},
            {"id": "3.2", "name": "Innovation Ecosystem & Startup Incubation"},
            {"id": "3.3", "name": "Research Publications, Books & Citations"},
            {"id": "3.4", "name": "Extension and Community Outreach Activities"},
            {"id": "3.5", "name": "Collaborations, MoUs & Linkages"}
        ]
    },
    {
        "id": "Criterion 4",
        "name": "Infrastructure and Learning Resources",
        "weightage": 100,
        "description": "Physical instructional facilities, modern central library automation, high-speed campus IT networks, and estate maintenance protocols.",
        "subcriteria": [
            {"id": "4.1", "name": "Physical Facilities & Smart Classrooms"},
            {"id": "4.2", "name": "Library as a Learning Resource & ILMS"},
            {"id": "4.3", "name": "IT Infrastructure, Bandwidth & Wi-Fi"},
            {"id": "4.4", "name": "Campus Infrastructure Maintenance"}
        ]
    },
    {
        "id": "Criterion 5",
        "name": "Student Support and Progression",
        "weightage": 140,
        "description": "Scholarships, proctorial mentoring, career placement training, competitive examination coaching, sports/cultural bodies, and active alumni chapters.",
        "subcriteria": [
            {"id": "5.1", "name": "Student Support & Scholarships"},
            {"id": "5.2", "name": "Student Progression & Career Placements"},
            {"id": "5.3", "name": "Student Participation in Sports & Culture"},
            {"id": "5.4", "name": "Alumni Engagement & Contribution"}
        ]
    },
    {
        "id": "Criterion 6",
        "name": "Governance, Leadership and Management",
        "weightage": 100,
        "description": "Institutional vision, strategic plan deployment, faculty empowerment initiatives (FDPs), transparent financial audits, and IQAC quality assurance.",
        "subcriteria": [
            {"id": "6.1", "name": "Institutional Vision and Participative Leadership"},
            {"id": "6.2", "name": "Strategy Development and Perspective Deployment"},
            {"id": "6.3", "name": "Faculty Empowerment Strategies & FDPs"},
            {"id": "6.4", "name": "Financial Management & Resource Mobilization"},
            {"id": "6.5", "name": "Internal Quality Assurance System (IQAC)"}
        ]
    },
    {
        "id": "Criterion 7",
        "name": "Institutional Values and Best Practices",
        "weightage": 100,
        "description": "Gender equity, rooftop solar power, water conservation, waste recycling, clean green campus, distinctiveness, and institutional best practices.",
        "subcriteria": [
            {"id": "7.1", "name": "Institutional Values & Environmental Responsibility"},
            {"id": "7.2", "name": "Best Practices in Institutional Excellence"},
            {"id": "7.3", "name": "Institutional Distinctiveness"}
        ]
    }
]

DEPARTMENTS = [
    {"name": "IQAC", "head": "Dr. Rameshwar Rao", "code": "IQAC"},
    {"name": "Research Cell", "head": "Dr. Meenakshi Sundaram", "code": "RCELL"},
    {"name": "Academic Office", "head": "Dr. K. S. Murthy", "code": "ACAD"},
    {"name": "Training & Placement Cell", "head": "Prof. S. R. V. Prasad", "code": "TPC"},
    {"name": "Student Affairs", "head": "Dr. V. Lalitha Devi", "code": "SAFF"},
    {"name": "Library", "head": "Dr. G. Satyanarayana", "code": "LIB"},
    {"name": "Examination Cell", "head": "Prof. B. Prabhakar", "code": "EXAM"},
    {"name": "Administration", "head": "Sri K. Raghavan, Registrar", "code": "ADMIN"},
    {"name": "CSE Department", "head": "Dr. T. Venkat Reddy", "code": "CSE"},
    {"name": "ECE Department", "head": "Dr. M. Sridevi", "code": "ECE"},
    {"name": "Mechanical Department", "head": "Dr. P. V. Ramana", "code": "MECH"},
    {"name": "Civil Department", "head": "Dr. Ch. Krishna", "code": "CIVIL"},
    {"name": "EEE Department", "head": "Dr. A. Sudhir", "code": "EEE"},
    {"name": "Estate Office", "head": "Er. D. Madhav Rao", "code": "EST"}
]


class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_elements(num_pages)
            super().showPage()
        super().save()

    def draw_page_elements(self, page_count):
        self.saveState()
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.HexColor("#4c1d95"))
        self.drawString(54, 750, f"{INSTITUTION_NAME} ({INSTITUTION_SHORT}) — Accreditation Evidence Archive")
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#6b7280"))
        self.drawRightString(558, 750, "Internal Quality Assurance Cell (IQAC)")
        self.setStrokeColor(colors.HexColor("#e5e7eb"))
        self.setLineWidth(0.75)
        self.line(54, 742, 558, 742)

        self.line(54, 45, 558, 45)
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#9ca3af"))
        self.drawString(54, 32, "Confidential — Institutional Accreditation Evidence Document")
        page_str = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 32, page_str)
        self.restoreState()

styles = getSampleStyleSheet()
title_style = ParagraphStyle(
    'DocTitle',
    parent=styles['Heading1'],
    fontName='Helvetica-Bold',
    fontSize=18,
    leading=22,
    textColor=colors.HexColor('#2e1065'),
    spaceAfter=6
)
subtitle_style = ParagraphStyle(
    'DocSubTitle',
    parent=styles['Normal'],
    fontName='Helvetica',
    fontSize=10,
    leading=14,
    textColor=colors.HexColor('#6b7280'),
    spaceAfter=12
)
meta_style = ParagraphStyle(
    'MetaBadge',
    parent=styles['Normal'],
    fontName='Helvetica-Bold',
    fontSize=8,
    leading=11,
    textColor=colors.HexColor('#4c1d95')
)
h1_style = ParagraphStyle(
    'SecH1',
    parent=styles['Heading2'],
    fontName='Helvetica-Bold',
    fontSize=13,
    leading=17,
    textColor=colors.HexColor('#3b0764'),
    spaceBefore=12,
    spaceAfter=6
)
body_style = ParagraphStyle(
    'SecBody',
    parent=styles['Normal'],
    fontName='Helvetica',
    fontSize=9,
    leading=14,
    textColor=colors.HexColor('#1f2937'),
    spaceAfter=8
)
quote_style = ParagraphStyle(
    'SecQuote',
    parent=styles['Normal'],
    fontName='Helvetica-Oblique',
    fontSize=8.5,
    leading=13,
    textColor=colors.HexColor('#374151'),
    leftIndent=15,
    rightIndent=15,
    spaceAfter=8
)

def build_pdf_document(doc_data):
    filepath = os.path.join(DOCS_DIR, doc_data["filename"])
    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )
    story = []

    # Title Banner
    story.append(Paragraph(doc_data["title"], title_style))
    story.append(Paragraph(
        f"Institution: {INSTITUTION_NAME} ({INSTITUTION_SHORT}) | Academic Year: <b>{doc_data['academic_year']}</b> | Department: <b>{doc_data['department']}</b>",
        subtitle_style
    ))

    # Meta Table
    meta_table_data = [
        [
            Paragraph(f"<b>Document ID:</b> {doc_data['document_id']}", meta_style),
            Paragraph(f"<b>Criterion:</b> {doc_data['criterion']}", meta_style),
            Paragraph(f"<b>Status:</b> {doc_data.get('status', 'active').upper()}", meta_style)
        ],
        [
            Paragraph(f"<b>Subcriterion:</b> {doc_data.get('subcriterion', 'General')}", meta_style),
            Paragraph(f"<b>Lead Officer:</b> {doc_data.get('author', 'Director IQAC')}", meta_style),
            Paragraph(f"<b>Approved Date:</b> {doc_data.get('effective_date', '2025-05-10')}", meta_style)
        ]
    ]
    t = Table(meta_table_data, colWidths=[170, 190, 144])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f5f3ff')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#ddd6fe')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#ede9fe')),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(t)
    story.append(Spacer(1, 12))

    # Executive Summary
    story.append(Paragraph("<b>EXECUTIVE SUMMARY & ACCREDITATION SCOPE</b>", h1_style))
    story.append(Paragraph(doc_data["summary"], body_style))
    story.append(Spacer(1, 8))

    # Sections with explicit page flow
    sections = doc_data.get("sections", [])
    current_page = 1
    for s_idx, sec in enumerate(sections):
        target_page = sec.get("page_number", current_page)
        while current_page < target_page:
            story.append(PageBreak())
            current_page += 1

        story.append(Paragraph(f"Section {s_idx + 1}: {sec['section_title']}", h1_style))
        for p in sec.get("paragraphs", []):
            story.append(Paragraph(p, body_style))

        if "table" in sec and sec["table"]:
            tbl_info = sec["table"]
            tbl_data = []
            # Header row
            header_row = [Paragraph(f"<b>{col}</b>", meta_style) for col in tbl_info["headers"]]
            tbl_data.append(header_row)
            for row in tbl_info["rows"]:
                row_cells = [Paragraph(str(cell), body_style) for cell in row]
                tbl_data.append(row_cells)

            col_w = 504 / len(tbl_info["headers"])
            tbl = Table(tbl_data, colWidths=[col_w] * len(tbl_info["headers"]))
            tbl.setStyle(TableStyle([
                ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#ede9fe')),
                ('BOX', (0,0), (-1,-1), 0.75, colors.HexColor('#c4b5fd')),
                ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e5e7eb')),
                ('TOPPADDING', (0,0), (-1,-1), 4),
                ('BOTTOMPADDING', (0,0), (-1,-1), 4),
                ('LEFTPADDING', (0,0), (-1,-1), 6),
                ('RIGHTPADDING', (0,0), (-1,-1), 6),
            ]))
            story.append(Spacer(1, 6))
            story.append(tbl)
            story.append(Spacer(1, 8))

    doc.build(story, canvasmaker=NumberedCanvas)
    return filepath

def build_docx_document(doc_data):
    docx_filename = doc_data["filename"].replace(".pdf", ".docx")
    filepath = os.path.join(DOCS_DIR, docx_filename)
    d = docx.Document()
    d.add_heading(f"{INSTITUTION_NAME} ({INSTITUTION_SHORT})", level=0)
    d.add_heading(doc_data["title"], level=1)
    d.add_paragraph(f"Academic Year: {doc_data['academic_year']} | Department: {doc_data['department']} | Criterion: {doc_data['criterion']}")
    d.add_heading("Executive Summary", level=2)
    d.add_paragraph(doc_data["summary"])

    for s_idx, sec in enumerate(doc_data.get("sections", [])):
        d.add_heading(f"Section {s_idx + 1}: {sec['section_title']} (Page {sec.get('page_number', 1)})", level=2)
        for p in sec.get("paragraphs", []):
            d.add_paragraph(p)
        if "table" in sec and sec["table"]:
            tbl_info = sec["table"]
            table = d.add_table(rows=1, cols=len(tbl_info["headers"]))
            hdr_cells = table.rows[0].cells
            for i, name in enumerate(tbl_info["headers"]):
                hdr_cells[i].text = name
            for row in tbl_info["rows"]:
                row_cells = table.add_row().cells
                for i, val in enumerate(row):
                    row_cells[i].text = str(val)

    d.save(filepath)
    return filepath

print("ReportLab and python-docx build pipelines configured.")
